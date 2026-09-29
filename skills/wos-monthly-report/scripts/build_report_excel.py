#!/usr/bin/env python3
"""Build the institution/month WOS report Excel (+CSV) from the three intermediate files.

Usage:
  python3 build_report_excel.py --alldb ALLDB.tsv --woscc WOSCC.tsv [--cas CAS.json] \\
      --org-en "Minnan Normal Univ" --institution 闽南师范大学 --month 2026年7月 \\
      --out OUT.xlsx [--manual manual.json] [--no-csv]

  --alldb   All-Databases Tab-delimited export (from wos_alldb_search_export.py)
  --woscc   Core-Collection Full Record export (from woscc_full_export.py) -> supplies C1
  --cas     raw partition JSON (from letpub_cas_fetch.py); omit to skip partition columns
  --org-en  the ROMANIZED institution token as WOS writes it (for 第一单位 matching)
  --manual  JSON {UT: true|false} overriding first-unit for records whose C1 was missing
            (judge those from /wos/woscc/full-record/<UT>, 'Addresses' -> address 1)

Output columns: 序号 | <institution>是否第一单位 | 篇名 | 作者 | 期刊来源 |
                [中科院大类学科 | 大类分区 | 小类分区 | Top期刊] |
                入库时间 | 正式出版时间 | EA时间     (+ a 说明 sheet)
Non-formally-published records (UT prefix PPRN:/RC:) are dropped and reported, never
silently. Parsing is BY COLUMN INDEX (the alldb header repeats AU -> DictReader loses authors).
"""
import argparse
import csv
import json
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Reuse the CAS parser from the cas-journal-partition skill.
CASPATH = os.path.expanduser("~/.hermes/skills/research/cas-journal-partition/scripts")
if CASPATH not in sys.path:
    sys.path.insert(0, CASPATH)
try:
    from letpub_cas_parse import find_current_block, parse_cas_block
except Exception:  # noqa: BLE001
    find_current_block = parse_cas_block = None

MON = {'JAN': '01', 'FEB': '02', 'MAR': '03', 'APR': '04', 'MAY': '05', 'JUN': '06',
       'JUL': '07', 'AUG': '08', 'SEP': '09', 'OCT': '10', 'NOV': '11', 'DEC': '12'}


def norm_date(s):
    """AUG 27 2026 / AUG 2026 / 2026-Aug-27 / 2026 -> ISO-ish YYYY-MM[-DD]."""
    s = (s or '').strip()
    if not s:
        return ''
    m = re.match(r'^([A-Za-z]{3})\s+(\d{1,2})\s+(\d{4})$', s)
    if m:
        return f"{m.group(3)}-{MON.get(m.group(1).upper(), '?')}-{int(m.group(2)):02d}"
    m = re.match(r'^([A-Za-z]{3})\s+(\d{4})$', s)
    if m:
        return f"{m.group(2)}-{MON.get(m.group(1).upper(), '?')}"
    m = re.match(r'^(\d{4})-([A-Za-z]{3})-(\d{1,2})$', s)
    if m:
        return f"{m.group(1)}-{MON.get(m.group(2).upper(), '?')}-{int(m.group(3)):02d}"
    # month-only variants: "2026-Jun" (dashed) — the day-less MEDLINE form
    m = re.match(r'^(\d{4})-([A-Za-z]{3})$', s)
    if m:
        return f"{m.group(1)}-{MON.get(m.group(2).upper(), '?')}"
    # MEDLINE rows can also arrive space-separated: "2026 Jun 11" / "2026 Jun"
    m = re.match(r'^(\d{4})\s+([A-Za-z]{3})\s+(\d{1,2})$', s)
    if m:
        return f"{m.group(1)}-{MON.get(m.group(2).upper(), '?')}-{int(m.group(3)):02d}"
    m = re.match(r'^(\d{4})\s+([A-Za-z]{3})$', s)
    if m:
        return f"{m.group(1)}-{MON.get(m.group(2).upper(), '?')}"
    # 双月刊/合刊期号："SEP-OCT 2026" / "NOV-DEC 2026" / "2026 SEP-OCT"（取区间首月）
    m = re.match(r'^([A-Za-z]{3})\s*[-/]\s*([A-Za-z]{3})\s+(\d{4})$', s)
    if m:
        return f"{m.group(3)}-{MON.get(m.group(1).upper(), '?')}"
    m = re.match(r'^(\d{4})\s+([A-Za-z]{3})\s*[-/]\s*([A-Za-z]{3})$', s)
    if m:
        return f"{m.group(1)}-{MON.get(m.group(2).upper(), '?')}"
    return s


def read_indexed(path):
    """Return (header, first_index_by_name, rows_as_lists)."""
    with open(path, encoding='utf-8-sig') as f:
        lines = f.readlines()
    header = lines[0].rstrip('\n').split('\t')
    idx = {}
    for i, n in enumerate(header):
        if n and n not in idx:
            idx[n] = i
    rows = []
    for line in lines[1:]:
        if not line.strip():
            continue
        c = line.rstrip('\n').split('\t')
        if len(c) < len(header):
            c += [''] * (len(header) - len(c))
        rows.append(c)
    return header, idx, rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--alldb', required=True)
    ap.add_argument('--woscc', required=True)
    ap.add_argument('--cas')
    ap.add_argument('--org-en', required=True)
    ap.add_argument('--institution', required=True)
    ap.add_argument('--month', required=True)
    ap.add_argument('--out', required=True)
    ap.add_argument('--manual')
    ap.add_argument('--no-csv', action='store_true')
    a = ap.parse_args()

    manual = {}
    if a.manual and os.path.exists(a.manual):
        manual = {k: (v if isinstance(v, str) else ('是' if v else '否'))
                  for k, v in json.load(open(a.manual, encoding='utf-8')).items()}

    _, ai, arows = read_indexed(a.alldb)
    keep, dropped = [], []
    for c in arows:
        ut = c[ai['UT']]
        (keep if ut.startswith(('WOS:', 'MEDLINE:')) else dropped).append(c)
    print(f"[in] alldb 记录 {len(arows)} -> 正式发表 {len(keep)}，剔除非期刊 {len(dropped)} "
          f"({', '.join(sorted({c[ai['UT']].split(':')[0] for c in dropped})) or '无'})")

    _, wi, wrows = read_indexed(a.woscc)
    c1 = {c[wi['UT']]: c[wi['C1']] for c in wrows}

    token = a.org_en.upper()

    def first_is_org(c1v):
        m = re.match(r'^\[([^\]]*)\]\s*(.*?)(?:;\s*\[|$)', c1v or '')
        return bool(m and token in (m.group(2) or '').upper())

    parsed = {}
    if a.cas and parse_cas_block:
        raw = json.load(open(a.cas, encoding='utf-8'))
        for jn, d in raw.items():
            if isinstance(d, dict) and 'error' not in d and d.get('block'):
                ver, blk = find_current_block('期刊分区表\n（ ' + d.get('version', '') + '）\n' + d['block'])
                r = parse_cas_block(blk) if blk else parse_cas_block(d['block'])
                if r:
                    parsed[jn] = {**r, 'version': d.get('version', '')}
        print(f"[cas] 分区数据 {len(parsed)} 刊")

    def first_unit(ut):
        if ut in manual:
            return manual[ut]
        if ut not in c1:
            return '待定'
        return '是' if first_is_org(c1[ut]) else '否'

    rows, pend = [], []
    for i, c in enumerate(keep, 1):
        ut = c[ai['UT']]
        so = c[ai['SO']]
        fu = first_unit(ut)
        if fu == '待定':
            pend.append(ut)
        pr = parsed.get(so, {})
        rec = {
            '序号': i,
            f'{a.institution}是否第一单位': fu,
            '篇名': c[ai['TI']], '作者': c[ai['AU']], '期刊来源': so,
            '入库时间': norm_date(c[ai['DA']]), '正式出版时间': norm_date(c[ai['PD']]),
            'EA时间': norm_date(c[ai['EA']]),
        }
        if a.cas:
            rec['中科院大类学科'] = pr.get('大类', '')
            rec['大类分区'] = f"{pr['大类区']}区" if pr.get('大类区') else ''
            rec['小类分区'] = '；'.join(f"{cn} {q}区" if q else cn for _, cn, q in pr.get('小类', []))
            rec['Top期刊'] = pr.get('Top', '')
        rows.append(rec)

    if pend:
        print(f"[warn] {len(pend)} 条缺 C1，第一单位=待定；逐条查 /wos/woscc/full-record/<UT> 的 "
              f"Addresses 第1条后用 --manual 重跑: {pend}")

    if a.cas:
        fields = ['序号', f'{a.institution}是否第一单位', '篇名', '作者', '期刊来源',
                  '中科院大类学科', '大类分区', '小类分区', 'Top期刊',
                  '入库时间', '正式出版时间', 'EA时间']
    else:
        fields = ['序号', f'{a.institution}是否第一单位', '篇名', '作者', '期刊来源',
                  '入库时间', '正式出版时间', 'EA时间']

    wb = Workbook(); ws = wb.active; ws.title = '论文列表'
    hf = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
    hfill = PatternFill('solid', start_color='1F4E79')
    hcent = Alignment(horizontal='center', vertical='center', wrap_text=True)
    bf = Font(name='微软雅黑', size=10)
    bwrap = Alignment(vertical='center', wrap_text=True)
    ccent = Alignment(horizontal='center', vertical='center')
    bd = Border(*[Side(style='thin', color='D9D9D9')] * 4)
    fy = PatternFill('solid', start_color='C6EFCE'); fn = PatternFill('solid', start_color='FFC7CE')
    center_cols = {'序号', f'{a.institution}是否第一单位', '中科院大类学科', '大类分区', 'Top期刊',
                   '入库时间', '正式出版时间', 'EA时间'}
    for c, n in enumerate(fields, 1):
        cell = ws.cell(1, c, n); cell.font = hf; cell.fill = hfill; cell.alignment = hcent; cell.border = bd
    ws.row_dimensions[1].height = 30
    for ri, r in enumerate(rows, 2):
        for c, n in enumerate(fields, 1):
            cell = ws.cell(ri, c, r[n]); cell.font = bf; cell.border = bd
            cell.alignment = ccent if n in center_cols else bwrap
        flag = ws.cell(ri, 2)
        if r[f'{a.institution}是否第一单位'] == '是':
            flag.fill = fy
        elif r[f'{a.institution}是否第一单位'] == '否':
            flag.fill = fn
    widths = {'序号': 5, f'{a.institution}是否第一单位': 16, '篇名': 52, '作者': 40, '期刊来源': 30,
              '中科院大类学科': 14, '大类分区': 9, '小类分区': 40, 'Top期刊': 8,
              '入库时间': 12, '正式出版时间': 13, 'EA时间': 11}
    for i, n in enumerate(fields, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths[n]
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f"A1:{get_column_letter(len(fields))}{len(rows) + 1}"

    from collections import Counter
    fu_c = Counter(r[f'{a.institution}是否第一单位'] for r in rows)
    pq_c = Counter(r.get('大类分区', '') for r in rows if r.get('大类分区'))
    notes = [
        ['检索对象', a.institution], ['月份', a.month],
        ['检索数据库', 'Web of Science 所有数据库 (All Databases)'],
        ['检索字段', f'Address = "{a.org_en}"'],
        ['原始条数 / 非期刊剔除 / 正式发表', f"{len(arows)} / {len(dropped)} / {len(rows)}"],
        ['第一单位判定', 'WOS记录C1地址列表第一组(第一作者单位)'],
        [f'{a.institution}为第一单位', f"{fu_c.get('是', 0)} 篇"],
        ['非第一单位 / 待定', f"{fu_c.get('否', 0)} / {fu_c.get('待定', 0)}"],
        ['日期格式', '已归一化为 YYYY-MM-DD / YYYY-MM'],
    ]
    if a.cas:
        notes += [['中科院分区来源', 'LetPub(转引中科院文献情报中心期刊分区表)'],
                  ['分区版本', next((v['version'] for v in parsed.values() if v.get('version')), '')],
                  ['大类分区分布', '；'.join(f'{k}:{v}篇' for k, v in sorted(pq_c.items()))],
                  ['Top期刊论文', f"{sum(1 for r in rows if r.get('Top期刊') == '是')} 篇"]]
    else:
        notes += [['中科院分区', '未包含(未提供 --cas)']]
    ws2 = wb.create_sheet('说明')
    for ri, row in enumerate(notes, 1):
        for c, v in enumerate(row, 1):
            ws2.cell(ri, c, v)
    ws2.column_dimensions['A'].width = 26; ws2.column_dimensions['B'].width = 74
    wb.save(a.out)
    print(f"[OK] {a.out} | 第一单位 {dict(fu_c)} | 分区 {dict(pq_c)}")

    if not a.no_csv:
        csvpath = os.path.splitext(a.out)[0] + '.csv'
        with open(csvpath, 'w', newline='', encoding='utf-8-sig') as f:
            w = csv.DictWriter(f, fieldnames=fields); w.writeheader()
            for r in rows:
                w.writerow(r)
        print(f"[OK] {csvpath}")


if __name__ == '__main__':
    main()
