#!/usr/bin/env python3
"""公开发布前的交付物脱敏：删除「作者」列、移除「第一作者贡献率」工作表，并核对结果。

用法: python3 deidentify_downloads.py <目录>
只改动传入目录下的副本（如 portal_publish/downloads），不动工作区原始交付物。
"""
import glob
import os
import sys

from openpyxl import load_workbook

root = sys.argv[1] if len(sys.argv) > 1 else "$WOS_WORKDIR/portal_publish/downloads"
changed = []

for p in sorted(glob.glob(f"{root}/**/*.xlsx", recursive=True)):
    wb = load_workbook(p)
    notes = []
    for ws in list(wb.worksheets):
        if "第一作者" in ws.title:                     # 含姓名的工作表整体移除
            wb.remove(ws)
            notes.append(f"移除工作表「{ws.title}」")
            continue
        r = next(ws.iter_rows(min_row=1, max_row=1, values_only=False), [])
        for c in r:
            if str(c.value or "").strip() in ("作者", "第一作者", "通信作者"):
                ws.delete_cols(c.column, 1)
                notes.append(f"「{ws.title}」删除列「{c.value}」")
                break
    if notes:
        wb.save(p)
        changed.append((p, notes))

print(f"处理目录: {root}")
for p, notes in changed:
    print(f"  ✓ {os.path.relpath(p, root)}")
    for n in notes:
        print(f"      - {n}")
if not changed:
    print("  无需改动")

print("\n=== 复核 ===")
for p in sorted(glob.glob(f"{root}/**/*.xlsx", recursive=True)):
    wb = load_workbook(p, read_only=True)
    bad_sheets = [s for s in wb.sheetnames if "第一作者" in s]
    hdrs = []
    for s in wb.sheetnames:
        ws = wb[s]
        hdrs += [str(c) for c in next(ws.iter_rows(max_row=1, values_only=True), ()) if c]
    bad_cols = [h for h in hdrs if h.strip() in ("作者", "第一作者", "通信作者")]
    flag = "OK" if not bad_sheets and not bad_cols else "仍有残留"
    print(f"  [{flag}] {os.path.relpath(p, root)} | 工作表={wb.sheetnames} | 姓名字段={bad_cols or '无'}")
